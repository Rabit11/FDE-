# -*- coding: utf-8 -*-
"""生成「制度管理」单页 PPT"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = r"C:\Users\81172\Desktop\制度管理.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

BLUE = RGBColor(0x00, 0x48, 0xA0)
BLUE2 = RGBColor(0x00, 0x64, 0xEF)
LIGHT = RGBColor(0xEE, 0xF4, 0xFC)
DARK = RGBColor(0x1F, 0x23, 0x29)
MUTED = RGBColor(0x4E, 0x59, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xE4, 0xF5)
BG = RGBColor(0xF5, 0xF7, 0xFA)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
ARROW = RGBColor(0x9B, 0xBF, 0xE8)


def set_fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def set_run(run, text, size, bold, color):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


# 背景
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
set_fill(bg, BG)

# 顶栏
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
set_fill(bar, BLUE)

# 标题
title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.9), Inches(0.7))
tp = title_box.text_frame.paragraphs[0]
tp.alignment = PP_ALIGN.LEFT
set_run(tp.add_run(), "制度管理", 36, True, BLUE)

uline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.15), Inches(1.2), Inches(0.06))
set_fill(uline, BLUE2)

# 标准原文
quote = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.45), Inches(11.9), Inches(1.4)
)
quote.fill.solid()
quote.fill.fore_color.rgb = WHITE
quote.line.color.rgb = LINE
quote.line.width = Pt(1)
qtf = quote.text_frame
qtf.word_wrap = True
qtf.margin_left = Inches(0.35)
qtf.margin_right = Inches(0.35)
qtf.margin_top = Inches(0.22)
qp = qtf.paragraphs[0]
set_run(qp.add_run(), "标准要求", 12, True, BLUE2)
qp2 = qtf.add_paragraph()
qp2.space_before = Pt(10)
set_run(
    qp2.add_run(),
    "班组能够结合管理实际，将上级组织的管理制度细化形成适合本班组的管理制度，"
    "并持续进行优化改进，提升制度的执行力。",
    16,
    False,
    DARK,
)

# 四步流程
cards = [
    ("01", "结合实际", ["对照上级管理制度", "结合班组管理实际"]),
    ("02", "细化落地", ["细化形成适合本班组的", "管理制度与细则"]),
    ("03", "持续优化", ["定期评估与改进", "动态完善制度体系"]),
    ("04", "提升执行力", ["强化宣贯与落实", "提高制度执行成效"]),
]

x0, gap, cw = 0.7, 0.28, 2.8
for i, (num, title, lines) in enumerate(cards):
    x = Inches(x0 + i * (cw + gap))
    y = Inches(3.2)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(cw), Inches(2.9))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    card.line.width = Pt(1)

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(cw), Inches(0.1))
    set_fill(top, BLUE if i < 3 else BLUE2)

    nb = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.95), y + Inches(0.35), Inches(0.9), Inches(0.9))
    set_fill(nb, LIGHT)
    ntf = nb.text_frame
    ntf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(ntf.paragraphs[0].add_run(), num, 18, True, BLUE)

    tb = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(1.4), Inches(cw - 0.3), Inches(0.45))
    ttf = tb.text_frame
    ttf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(ttf.paragraphs[0].add_run(), title, 20, True, BLUE)

    db = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(1.95), Inches(cw - 0.4), Inches(0.8))
    dtf = db.text_frame
    dtf.word_wrap = True
    for j, line in enumerate(lines):
        pp = dtf.paragraphs[0] if j == 0 else dtf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        set_run(pp.add_run(), line, 13, False, MUTED)

    if i < 3:
        ar = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            x + Inches(cw) + Inches(0.02),
            y + Inches(1.25),
            Inches(0.24),
            Inches(0.28),
        )
        set_fill(ar, ARROW)

# 页脚
fb = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(11.9), Inches(0.35))
fp = fb.text_frame.paragraphs[0]
set_run(fp.add_run(), "班组建设 · 制度管理", 11, False, GRAY)

prs.save(OUT)
print("saved", OUT)
